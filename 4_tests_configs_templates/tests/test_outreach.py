"""Tests for measure_it.outreach (generic outreach workbook)."""
from __future__ import annotations

import re

import pandas as pd
import pytest

from measure_it import outreach as O

BANNED = re.compile(r"\bbest clinic|\bproves?\b|diagnosed by ai|patient has|definitive biomarker|optimal treatment",
                    re.IGNORECASE)


def _docs(n, county="00001", **kw):
    base = {"npi": [f"{i:010d}" for i in range(n)], "county_fips": [county] * n, "act_score": [0] * n,
            "act_services": [0] * n, "rx_relation_score": [0] * n, "rx_claims": [0] * n, "exp_score": [0] * n,
            "specialty_groups": ["primary_care"] * n, "last_update_date": pd.to_datetime(["2020-01-01"] * n),
            "address_key": [f"A{i}|00001" for i in range(n)]}
    base.update(kw)
    return pd.DataFrame(base)


def test_norm_address():
    assert O.norm_address("100 Main St., Suite 200") == "100 MAIN ST"
    assert O.norm_address("100 MAIN ST STE 3") == "100 MAIN ST"
    assert O.norm_address(None) == ""


def test_rank_order_activity_rx_site_specialty_recency():
    d = _docs(6, act_score=[0, 3, 1, 0, 0, 0], act_services=[0, 5, 50, 0, 0, 0], rx_relation_score=[0, 0, 0, 2, 0, 0],
              rx_claims=[0, 0, 0, 30, 0, 0], exp_score=[0, 0, 0, 0, 2, 0],
              specialty_groups=["primary_care", "primary_care", "primary_care", "primary_care", "primary_care",
                                "rheumatology"],
              last_update_date=pd.to_datetime(["2025-01-01"] + ["2020-01-01"] * 5))
    r = O.rank_clinicians(d, {"primary_care": 1, "rheumatology": 2}, per_county=10)
    # dedicated (3) beats proxy (1) even with fewer services; then rx; then site; then specialist tier; then recency
    assert r["npi"].tolist() == ["0000000001", "0000000002", "0000000003", "0000000004", "0000000005", "0000000000"]
    assert r["priority"].tolist() == [1, 2, 3, 4, 5, 6]


def test_weak_proxy_activity_ranks_after_specialty():
    d = _docs(3, act_weak=[1, 0, 0], specialty_groups=["primary_care", "rheumatology", "primary_care"],
              last_update_date=pd.to_datetime(["2020-01-01", "2020-01-01", "2025-01-01"]))
    r = O.rank_clinicians(d, {"primary_care": 1, "rheumatology": 2}, per_county=10)
    # specialist first; then the weak-proxy biller ahead of a more recently updated primary-care record
    assert r["npi"].tolist() == ["0000000001", "0000000000", "0000000002"]


def test_rank_caps_per_address_and_county_and_drops_irrelevant():
    d = _docs(8, address_key=["X|1"] * 7 + ["Y|1"], specialty_groups=["primary_care"] * 7 + ["dermatology"])
    r = O.rank_clinicians(d, {"primary_care": 1}, per_county=20, max_per_address=5)
    assert len(r) == 5 and set(r["address_key"]) == {"X|1"}  # dermatology without evidence is not listed
    r2 = O.rank_clinicians(_docs(30), {"primary_care": 1}, per_county=20)
    assert len(r2) == 20


def test_resolve_condition_forms():
    assert O.resolve_condition("long_covid_or_me_cfs")["members"] == ["long_covid", "me_cfs"]
    s = O.resolve_condition("ptlds_or_lyme_disease")
    assert s["condition_id"] == "lyme_disease+ptlds" and s["members"] == ["lyme_disease", "ptlds"]
    assert O.resolve_condition("lyme_disease")["kind"] == "condition"


def test_specialty_tiers():
    cond = O.resolve_condition("long_covid")
    meas = O.resolve_measurement("nailfold_capillaroscopy")
    t = O.specialty_tiers(cond, meas)
    assert t["rheumatology"] == 3 and t["cardiology"] == 2 and t["primary_care"] == 1 and t["pulmonology"] == 1


@pytest.mark.data
def test_select_counties_scored_and_fallback():
    c, b = O.select_counties(O.resolve_condition("long_covid_or_me_cfs"),
                             O.resolve_measurement("nailfold_capillaroscopy"), 10)
    assert b["basis"] == "deployment_opportunities" and c["rank"].tolist() == list(range(1, 11))
    c2, b2 = O.select_counties(O.resolve_condition("ptlds_or_lyme_disease"),
                               O.resolve_measurement("nailfold_capillaroscopy"), 10)
    assert b2["basis"] in {"geo_subgroup_burden", "geo_condition_burden"} and len(c2) == 10
    assert (c2["population_total"] >= O.min_population()).all()
    assert "BURDEN-ONLY" in b2["description"] or b2["basis"] == "geo_subgroup_burden"


@pytest.mark.data
def test_build_outreach_workbook(tmp_path):
    from openpyxl import load_workbook
    res = O.build_outreach("long_covid_or_me_cfs", "autonomic_function_testing", top_n_counties=2, per_county=5,
                           out_dir=tmp_path)
    wb = load_workbook(res["path"])
    assert wb.sheetnames == ["README", "Counties", "Clinicians", "Research & experience sites"]
    readme = " ".join(str(c.value or "") for c in wb["README"]["A"])
    assert "does NOT mean the clinician sees Long COVID" in readme and not BANNED.search(readme)
    ws = wb["Clinicians"]
    hdr = [c.value for c in ws[1]]
    for col in ("Measurement activity (codes:services)", "Part D signals", "Experience-site match", "NPI registry",
                "NPPES last updated", "Outreach status"):
        assert col in hdr
    assert ws.max_row - 1 == res["clinicians_listed"] <= 10
    assert ws.data_validations.dataValidation  # status drop-down
    assert ws.cell(2, hdr.index("Name") + 1).font.name == "Arial"


@pytest.mark.data
def test_subgroup_fallback_when_requested():
    from measure_it.store import table_exists
    if not table_exists("geo_subgroup_burden"):
        pytest.skip("stage E table not built")
    c, b = O.select_counties(O.resolve_condition("long_covid"), O.resolve_measurement("nailfold_capillaroscopy"), 5,
                             prefer_subgroup=True)
    assert b["basis"] == "geo_subgroup_burden" and "BURDEN-ONLY" in b["description"] and len(c) == 5
    assert c["value"].is_monotonic_decreasing
