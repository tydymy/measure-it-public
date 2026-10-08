"""Generic outreach workbook: candidate counties x candidate clinicians for one condition (or set) x measurement.

Successor of the one-off ``outreach/build_outreach.py``; generic over condition ids / sets and measurement classes /
bundles (stage F, docs/HARMONIZATION_CONTRACT.md section 5).

Counties
    1. ``deployment_opportunities`` rank under the chosen weight set (the scored grid: long_covid, me_cfs, pots,
       dysautonomia and their sets x the four bundles).
    2. Fallback when the combination is not in the scored grid (e.g. ``lyme_disease``, ``ptlds``, a measurement
       class), or on request (``prefer_subgroup=True`` / ``--prefer-subgroup``): ``geo_subgroup_burden`` (stage E,
       contract section 4; the estimated number of adults in the device-defined subgroup, ``subgroup_estimate``) when
       every member condition has county rows (``phenotype_id`` picks the phenotype), else the primary burden measure
       of ``geo_condition_burden`` (e.g. CDC Lyme incidence per 100k). Counties with population >= scoring.yaml
       ``ranking.min_population``; a set is ranked by the mean of its members' burden percentiles over counties where
       every member with a burden has a value (the scoring's completeness rule). These are BURDEN-ONLY rankings,
       labelled as such in the workbook; they are not the deployment composite.

Clinicians (per county; NPPES individual NPIs, active, practice ZIP in the county; MD/DO only unless
``include_nonphysicians``; pediatric/adolescent taxonomies excluded unless ``include_pediatric``), ranked
lexicographically by
    1. measurement-specific activity: Medicare FFS billing of the measurement's codes (dedicated > related code of
       medium/high confidence, then services; ``provider_measurement_activity``). Low-confidence related and proxy
       codes are displayed but rank only after step 4 (they are not measurement-specific),
    2. condition prescribing signals (``provider_condition_rx_signals``; direct > related phenotype, then claims;
       >= 2 signal drugs or one of medium specificity; a single low-specificity drug ranks after step 4),
    3. experience-site affiliation (practice address = a measurement experience site's address; else same ZIP5 as a
       measurement experience site or a condition research site: co-location, not affiliation),
    4. implementer specialty (configs/relevance.yaml measurement_implementers of the member classes; primary care
       and the condition's core specialties rank below the measurement's specialist implementer groups),
    5. recency of the NPPES record,
with at most ``max_per_address`` clinicians per practice address. Clinicians without any evidence and without a
relevant specialty are not listed.

Named clinicians are written ONLY to ``out_dir`` (default ``outreach/``, untracked; never commit).

CLI: ``uv run python -m measure_it.outreach --condition long_covid_or_me_cfs --measurement nailfold_capillaroscopy
--top 10 --per-county 20``
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd

from .config import PROJECT_ROOT, UNKNOWN, load_config
from .store import read_table, table_exists

PHYS = "Allopathic & Osteopathic Physicians"
APP = "Physician Assistants & Advanced Practice Nursing Providers"
ROLE_SCORE = {"dedicated": 3, "related": 2, "proxy": 1, "none": 0}
REL_SCORE = {"direct": 2, "related_phenotype": 1}
STATUS = ("Not contacted,Emailed,Called,Left message,Replied,Interested,Meeting set,Declined,Wrong contact,"
          "Not relevant")
SEP = "|"


# ------------------------------------------------------------------------------------------------------------------
# resolution
# ------------------------------------------------------------------------------------------------------------------

def resolve_condition(condition_id: str) -> dict:
    """Condition id, scored set id, ad hoc 'a+b', or 'a_or_b' of canonical ids -> condition_spec."""
    from .scoring.opportunity import condition_spec
    try:
        return condition_spec(condition_id)
    except KeyError:
        if "_or_" in condition_id:
            return condition_spec("+".join(sorted(condition_id.split("_or_"))))
        raise


def resolve_measurement(measurement_id: str) -> dict:
    from .facilities import matching as M
    m = M.resolve_measurement(measurement_id)
    if m["status"] != "matched":
        raise KeyError(f"measurement {measurement_id!r}: {m.get('reason')}")
    return m


def specialty_tiers(cond: dict, meas: dict) -> dict[str, int]:
    """group -> score: 3 = specialist implementer group of the measurement's PRIMARY member class (the first member of
    a bundle in configs/relevance.yaml, e.g. capillaroscopy -> rheumatology, vascular), 2 = specialist implementer of
    another member class, 1 = primary care / the condition's core specialties. Curated assumptions."""
    rel = load_config("relevance")
    core = set()
    for c in cond["members"]:
        core |= set((rel["condition_specialties"].get(c) or {}).get("core", []))
    impl_cfg = rel.get("measurement_implementers") or {}
    primary = set((impl_cfg.get(meas["members"][0]) or {}).get("groups", [])) - {"primary_care"}
    impl = set(meas["implementer_groups"]) - {"primary_care"}
    tiers = {g: 1 for g in core | {"primary_care"}}
    tiers.update({g: 2 for g in impl})
    tiers.update({g: 3 for g in primary})
    return tiers


# ------------------------------------------------------------------------------------------------------------------
# counties
# ------------------------------------------------------------------------------------------------------------------

def min_population() -> int:
    return int(load_config("scoring").get("ranking", {}).get("min_population", 10000))


def _county_frame() -> pd.DataFrame:
    f = read_table("geo_condition_features", columns=["geo_id", "geo_level", "geo_name", "state_abbr",
                                                      "population_total", "condition_id", "in_analysis_universe"])
    f = f[(f["condition_id"] == "long_covid") & (f["geo_level"] == "county") & f["in_analysis_universe"]]
    return f.drop(columns=["condition_id", "in_analysis_universe", "geo_level"]).rename(
        columns={"geo_id": "county_fips"})


SUBGROUP_COUNT_COLS = ("subgroup_estimate", "estimated_adults", "estimated_adults_in_subgroup", "subgroup_adults")
SUBGROUP_RATE_COLS = ("subgroup_rate_per_100k", "estimated_rate", "subgroup_rate")


def _subgroup_burden(members: list[str], phenotype_id: str | None = None) -> tuple[pd.DataFrame | None, str]:
    """County subgroup estimates from stage E (contract section 4): the estimated number of adults in the subgroup
    ('where are many such people'), summed over member conditions. Used only when EVERY member has rows (one
    phenotype; `phenotype_id` picks it, else the first in sort order, named in the description)."""
    if not table_exists("geo_subgroup_burden"):
        return None, "geo_subgroup_burden not present"
    s = read_table("geo_subgroup_burden")
    s = s[(s["geo_level"] == "county") & s["condition_id"].isin(members)]
    if phenotype_id:
        s = s[s["phenotype_id"] == phenotype_id]
    if s.empty or set(s["condition_id"]) != set(members):
        return None, "geo_subgroup_burden has no county rows for every member condition"
    pid = sorted(s["phenotype_id"].unique())[0]
    s = s[s["phenotype_id"] == pid]
    col = next((c for c in SUBGROUP_COUNT_COLS + SUBGROUP_RATE_COLS if c in s.columns), None)
    if col is None:
        return None, "geo_subgroup_burden has no recognised estimate column"
    v = s.groupby("geo_id")[col].sum() if col in SUBGROUP_COUNT_COLS else s.groupby("geo_id")[col].mean()
    synthetic = " (SYNTHETIC phenotype: test data, not a real estimate)" if "synthetic" in pid else ""
    return v.rename("value").reset_index().rename(columns={"geo_id": "county_fips"}), \
        f"geo_subgroup_burden.{col} for phenotype {pid}{synthetic} (stage E model-based estimate)"


def _primary_burden(members: list[str]) -> tuple[pd.DataFrame | None, str, list[str]]:
    b = read_table("geo_condition_burden", columns=["geo_id", "geo_level", "condition_id", "measure_id",
                                                    "measure_label", "value", "is_primary_measure", "year",
                                                    "burden_evidence_level"])
    b = b[(b["geo_level"] == "county") & b["is_primary_measure"].astype(bool) & b["condition_id"].isin(members)]
    if b.empty:
        return None, "no primary county burden measure for these conditions", []
    b = b.sort_values("year").drop_duplicates(["geo_id", "condition_id"], keep="last")
    wide = b.pivot(index="geo_id", columns="condition_id", values="value")
    have = list(wide.columns)
    labels = b.drop_duplicates("condition_id").set_index("condition_id")
    desc = "; ".join(f"{c}: {labels.loc[c, 'measure_label']} ({labels.loc[c, 'measure_id']}, "
                     f"{int(labels.loc[c, 'year']) if pd.notna(labels.loc[c, 'year']) else '?'}, evidence level "
                     f"{labels.loc[c, 'burden_evidence_level']})" for c in have)
    return wide.reset_index().rename(columns={"geo_id": "county_fips"}), desc, have


def select_counties(cond: dict, meas: dict, top_n: int, weight_set: str = "equal", phenotype_id: str | None = None,
                    prefer_subgroup: bool = False) -> tuple[pd.DataFrame, dict]:
    """Top-N counties and a dict describing the basis of the selection."""
    rank_col = f"rank_{weight_set}"
    d = read_table("deployment_opportunities")
    d = d[(d["condition_id"] == cond["condition_id"]) & (d["measurement_id"] == meas["measurement_id"])
          & (d["geo_level"] == "county")]
    if len(d) and rank_col in d and d[rank_col].notna().any() and not prefer_subgroup:
        keep = ["geo_id", "geo_name", "state_abbr", "population_total", rank_col, "composite_" + weight_set,
                "burden_pct", "vulnerability_pct", "diagnostic_desert_pct", "clinic_capacity_pct",
                "research_readiness_pct", "p_top10_mc", "rank_mc_p05", "rank_mc_p95"]
        keep = [c for c in keep if c in d]
        c = d[d[rank_col] <= top_n].sort_values(rank_col)[keep].rename(
            columns={"geo_id": "county_fips", rank_col: "rank"})
        return c.reset_index(drop=True), {
            "basis": "deployment_opportunities",
            "description": f"rank {rank_col} of the scored deployment opportunities (condition {cond['condition_id']}"
                           f" x measurement {meas['measurement_id']}, county level; candidate deployment "
                           "opportunities, not a validated ranking)"}
    geo = _county_frame()
    geo = geo[geo["population_total"].fillna(0) >= min_population()]
    sub, sub_desc = _subgroup_burden(cond["members"], phenotype_id)
    if sub is not None:
        m = geo.merge(sub, on="county_fips", how="inner").dropna(subset=["value"])
        m["burden_score"] = m["value"]
        basis = {"basis": "geo_subgroup_burden",
                 "description": f"SUBGROUP-BURDEN-ONLY ranking ({sub_desc}); not the deployment composite"}
    else:
        wide, desc, have = _primary_burden(cond["members"])
        if wide is None:
            raise KeyError(f"no county ranking available for {cond['condition_id']}: not in deployment_opportunities,"
                           f" {sub_desc}, {desc}")
        m = geo.merge(wide, on="county_fips", how="inner")
        m = m.dropna(subset=have)  # completeness: every member with a burden must have a value
        pr = np.vstack([m[c].rank(pct=True, method="average").to_numpy() for c in have])
        m["burden_score"] = pr.mean(axis=0)
        for c in have:
            m[f"burden_{c}"] = m[c]
        basis = {"basis": "geo_condition_burden",
                 "description": f"BURDEN-ONLY fallback ranking (combination not in deployment_opportunities; "
                                f"{sub_desc}): primary county burden measure per member: {desc}; set = mean of member "
                                "percentiles over complete counties"}
    m = m.sort_values(["burden_score", "county_fips"], ascending=[False, True]).reset_index(drop=True)
    m["rank"] = np.arange(1, len(m) + 1)
    basis["universe"] = f"{len(m):,} counties with population >= {min_population():,} and a value"
    cols = ["rank", "county_fips", "geo_name", "state_abbr", "population_total", "burden_score"] + \
        [c for c in m.columns if c.startswith("burden_") and c != "burden_score"] + \
        (["value"] if "value" in m else [])
    return m[m["rank"] <= top_n][cols].reset_index(drop=True), basis


# ------------------------------------------------------------------------------------------------------------------
# evidence per clinician
# ------------------------------------------------------------------------------------------------------------------

def norm_address(a) -> str:
    a = re.sub(r"[^A-Z0-9 ]", " ", str(a or "").upper())
    a = re.split(r"\b(?:STE|SUITE|UNIT|FL|FLOOR|RM|ROOM|BLDG|APT)\b", a)[0]
    return re.sub(r"\s+", " ", a).strip()


def activity_evidence(measurement_id: str) -> pd.DataFrame:
    """npi -> best code role, services, codes billed for the measurement (provider_measurement_summary or computed)."""
    if table_exists("provider_measurement_summary"):
        s = read_table("provider_measurement_summary")
        s = s[s["measurement_id"] == measurement_id]
        if len(s):
            return s
    from .facilities.activity import provider_summary
    return provider_summary(measurement_id)


def activity_strength(members: list[str]) -> pd.DataFrame:
    """npi -> act_score (3 dedicated code, 2 related code of medium/high confidence, 0 otherwise), services on those
    codes, and act_weak (1 = only low-confidence related or proxy codes). Low-confidence/proxy billing is shown but
    ranks after specialty: it is not measurement-specific (e.g. limb arterial studies for capillaroscopy)."""
    cols = ["npi", "act_score", "act_services", "act_weak"]
    if not table_exists("provider_measurement_activity"):
        return pd.DataFrame(columns=cols)
    a = read_table("provider_measurement_activity", columns=["npi", "measurement_class", "hcpcs_cd", "code_role",
                                                             "code_confidence", "tot_services"])
    a = a[a["measurement_class"].isin(members)].drop_duplicates(["npi", "hcpcs_cd", "code_role"])
    if a.empty:
        return pd.DataFrame(columns=cols)
    a["score"] = np.select([a["code_role"] == "dedicated",
                            (a["code_role"] == "related") & a["code_confidence"].isin(["medium", "high"])], [3, 2], 0)
    a["spec_serv"] = np.where(a["score"] > 0, a["tot_services"], 0.0)
    g = a.groupby("npi")
    out = pd.DataFrame({"act_score": g["score"].max(), "act_services": g["spec_serv"].sum()})
    out["act_weak"] = (out["act_score"] == 0).astype(int)
    return out.reset_index()


def rx_evidence(members: list[str]) -> pd.DataFrame:
    """npi -> Part D signal strength for the condition (set). Strong (rx_relation_score 2 direct / 1 related phenotype)
    needs >= 2 distinct signal drugs or one drug of medium/high specificity; a single low-specificity drug (e.g.
    midodrine alone, mostly dialysis / orthostatic hypotension in older adults) is shown but is `rx_weak` and ranks
    only after specialty."""
    cols = ["npi", "rx_relation_score", "rx_claims", "rx_signals", "rx_weak"]
    if not table_exists("provider_condition_rx_signals"):
        return pd.DataFrame(columns=cols)
    r = read_table("provider_condition_rx_signals", columns=["npi", "condition_relation", "drug_generic",
                                                             "specificity", "tot_claims", "prescriber_type"])
    rel = r["condition_relation"].map(json.loads)
    r["relation"] = [max((REL_SCORE.get(v.get(m), 0) for m in members), default=0) for v in rel]
    r = r[r["relation"] > 0]
    if r.empty:
        return pd.DataFrame(columns=cols)
    r["lab"] = r["drug_generic"] + ": " + r["tot_claims"].astype(int).astype(str) + " claims" + \
        np.where(r["relation"] == 1, " (related phenotype)", "")
    g = r.groupby("npi")
    out = pd.DataFrame({"rel": g["relation"].max(), "rx_claims": g["tot_claims"].sum(),
                        "rx_signals": g["lab"].agg("; ".join), "n_drugs": g["drug_generic"].nunique(),
                        "spec": g["specificity"].agg(lambda x: bool(set(x) & {"medium", "high"}))})
    strong = (out["n_drugs"] >= 2) | out["spec"]
    out["rx_relation_score"] = np.where(strong, out["rel"], 0)
    out["rx_weak"] = (~strong).astype(int)
    out["rx_claims"] = np.where(strong, out["rx_claims"], 0)
    out["rx_signals"] = out["rx_signals"] + np.where(strong, "", " [single low-specificity drug: weak]")
    return out[["rx_relation_score", "rx_claims", "rx_signals", "rx_weak"]].reset_index()


def site_keys(measurement_id: str, members: list[str]) -> tuple[dict, dict]:
    """(address+zip -> site name) for measurement experience sites; (zip -> site names) for those and condition
    research sites (precision view)."""
    addr, zips = {}, {}
    if table_exists("measurement_experience_sites"):
        e = read_table("measurement_experience_sites")
        e = e[e["measurement_id"] == measurement_id]
        for r in e.itertuples():
            lab = f"{r.facility_name} [{r.evidence_basis}: {r.n_trials} trials, {r.n_nih_projects} NIH; " \
                  f"{r.member_classes_matched}]"
            if pd.notna(r.address) and pd.notna(r.zip5):
                addr.setdefault((norm_address(r.address), str(r.zip5)), lab)
            if pd.notna(r.zip5):
                zips.setdefault(str(r.zip5), []).append(f"measurement: {r.facility_name}")
    rs = read_table("research_site_registry", columns=["facility_name", "zip5", "condition_ids_literal",
                                                       "nih_condition_ids_precision"])
    hit = rs["condition_ids_literal"].fillna("").str.split(SEP).map(lambda x: bool(set(x) & set(members))) | \
        rs["nih_condition_ids_precision"].fillna("").str.split(SEP).map(lambda x: bool(set(x) & set(members)))
    for r in rs[hit & rs["zip5"].notna()].itertuples():
        zips.setdefault(str(r.zip5), []).append(f"condition research: {r.facility_name}")
    return addr, zips


def candidate_clinicians(county_fips: list[str], include_nonphysicians: bool = False,
                         include_pediatric: bool = False) -> pd.DataFrame:
    import duckdb
    from .store import processed_path
    groupings = [PHYS] + ([APP] if include_nonphysicians else [])
    con = duckdb.connect()
    df = con.execute(f"""
      select npi, provider_first_name, provider_last_name, name, credential, primary_taxonomy_display_name as taxonomy,
             primary_taxonomy_grouping, specialty_groups, address_line1, city, state, zip5, phone, last_update_date,
             county_fips
      from '{processed_path("providers")}'
      where entity_type = 1 and deactivation_date is null and county_fips in (select unnest(?))
        and primary_taxonomy_grouping in (select unnest(?))""", [county_fips, groupings]).df()
    con.close()
    if not include_pediatric:  # adult cohorts (MAESTRO, Long COVID / ME/CFS / Lyme adult studies)
        df = df[~df["taxonomy"].fillna("").str.contains(r"Pediatric|Neonat|Adolescent", regex=True)]
    return df


def rank_clinicians(docs: pd.DataFrame, tiers: dict[str, int], per_county: int, max_per_address: int = 5,
                    specialty_before_rx: bool = False) -> pd.DataFrame:
    """Pure ranking: needs columns county_fips, act_score, act_services, rx_relation_score, rx_claims, exp_score,
    specialty_groups, last_update_date, address_key (optional act_weak, rx_weak: weak evidence, ranked after
    specialty). Returns the listed rows with `priority`."""
    d = docs.copy()
    d["spec_score"] = d["specialty_groups"].fillna("").str.split(SEP).map(
        lambda gs: max((tiers.get(g, 0) for g in gs if g), default=0))
    for c in ("act_weak", "rx_weak"):
        if c not in d:
            d[c] = 0
    for c in ("act_score", "act_services", "rx_relation_score", "rx_claims", "exp_score", "act_weak", "rx_weak"):
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0)
    d["weak_evidence"] = d["act_weak"] + d["rx_weak"]
    d = d[(d["act_score"] > 0) | (d["rx_relation_score"] > 0) | (d["exp_score"] > 0) | (d["spec_score"] > 0)
          | (d["weak_evidence"] > 0)]
    d["_upd"] = pd.to_datetime(d["last_update_date"], errors="coerce")
    keys = ["act_score", "act_services", "rx_relation_score", "rx_claims", "exp_score", "spec_score"]
    if specialty_before_rx:  # measurement specialists (e.g. rheumatology for capillaroscopy) ahead of rx / sites
        keys = ["act_score", "act_services", "spec_score", "rx_relation_score", "rx_claims", "exp_score"]
    d = d.sort_values(["county_fips"] + keys + ["weak_evidence", "_upd", "npi"],
                      ascending=[True] + [False] * (len(keys) + 2) + [True], kind="stable")
    d = d[d.groupby(["county_fips", "address_key"]).cumcount() < max_per_address]
    d = d[d.groupby("county_fips").cumcount() < per_county].copy()
    d["priority"] = d.groupby("county_fips").cumcount() + 1
    return d.drop(columns="_upd")


# ------------------------------------------------------------------------------------------------------------------
# workbook
# ------------------------------------------------------------------------------------------------------------------

def _fmt_phone(p) -> str:
    return f"({p[:3]}) {p[3:6]}-{p[6:10]}" if isinstance(p, str) and len(p) >= 10 and p[:10].isdigit() else (p or "")


def _write_workbook(path: Path, readme: list[str], counties: pd.DataFrame, clin: pd.DataFrame,
                    sites: pd.DataFrame) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    H, HF = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    B, LINK = Font(name=F, size=10), Font(name=F, size=10, color="0563C1", underline="single")
    FILL_IN = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()

    def sheet(ws, df: pd.DataFrame, widths: dict | None = None, input_cols=()):
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font, cell.fill = H, HF
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) or v is pd.NA or v is pd.NaT else v
                       for v in row])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = B
                if ws.cell(1, cell.column).value in input_cols:
                    cell.fill = FILL_IN
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, min(45, max(10, len(col) + 2)))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws0 = wb.active
    ws0.title = "README"
    for i, t in enumerate(readme, 1):
        c = ws0.cell(i, 1, t)
        c.font = Font(name=F, bold=(i == 1 or (t.isupper() and len(t) < 40)), size=13 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 150

    ws1 = wb.create_sheet("Counties")
    sheet(ws1, counties, {"County": 30})
    ws2 = wb.create_sheet("Clinicians")
    inputs = ("Outreach status", "Date contacted", "Contact person / email", "Notes")
    sheet(ws2, clin, {"Name": 26, "Practice address": 32, "Measurement activity (codes:services)": 40,
                      "Part D signals": 40, "Experience-site match": 45, "NPPES specialty": 32, "County": 26},
          input_cols=inputs)
    cols = list(clin.columns)
    if len(clin):
        npi_c, link_c = cols.index("NPI") + 1, cols.index("NPI registry") + 1
        for i in range(2, ws2.max_row + 1):
            ws2.cell(i, link_c).hyperlink = f"https://npiregistry.cms.hhs.gov/provider-view/{ws2.cell(i, npi_c).value}"
            ws2.cell(i, link_c).font = LINK
            for name in ("NPPES last updated", "Date contacted"):
                ws2.cell(i, cols.index(name) + 1).number_format = "yyyy-mm-dd"
        L = get_column_letter(cols.index("Outreach status") + 1)
        dv = DataValidation(type="list", formula1=f'"{STATUS}"', allow_blank=True)
        ws2.add_data_validation(dv)
        dv.add(f"{L}2:{L}{max(ws2.max_row, 2)}")
    ws3 = wb.create_sheet("Research & experience sites")
    sheet(ws3, sites, {"Facility": 45, "Trials (NCT)": 40, "NIH projects (appl_id)": 30, "NIH contact PIs": 30})
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------------------------------------------------------
# entry point
# ------------------------------------------------------------------------------------------------------------------

def build_outreach(condition_id: str, measurement_id: str, top_n_counties: int = 10, per_county: int = 20,
                   weight_set: str = "equal", out_dir: str | Path = "outreach", include_nonphysicians: bool = False,
                   max_per_address: int = 5, include_pediatric: bool = False,
                   specialty_before_rx: bool = False, phenotype_id: str | None = None,
                   prefer_subgroup: bool = False) -> dict:
    cond = resolve_condition(condition_id)
    meas = resolve_measurement(measurement_id)
    mid = meas["measurement_id"]
    counties, basis = select_counties(cond, meas, top_n_counties, weight_set, phenotype_id, prefer_subgroup)
    fips = counties["county_fips"].tolist()

    docs = candidate_clinicians(fips, include_nonphysicians, include_pediatric)
    act = activity_evidence(mid)
    act = act[["npi", "best_code_role", "tot_services", "codes_billed", "cms_provider_type", "county_fips"]].rename(
        columns={"tot_services": "act_services", "county_fips": "act_county"})
    docs = docs.merge(act.drop(columns=["act_county", "act_services"]), on="npi", how="left")
    docs = docs.merge(activity_strength(meas["members"]), on="npi", how="left")
    docs["act_services_all"] = docs["npi"].map(act.set_index("npi")["act_services"])
    rx = rx_evidence(cond["members"])
    docs = docs.merge(rx, on="npi", how="left")
    addr, zips = site_keys(mid, cond["members"])
    docs["address_key"] = [norm_address(a) + "|" + str(z) for a, z in zip(docs["address_line1"], docs["zip5"])]
    exp_addr = [addr.get((norm_address(a), str(z)), "") for a, z in zip(docs["address_line1"], docs["zip5"])]
    exp_zip = ["; ".join(sorted(set(zips.get(str(z), [])))[:3]) for z in docs["zip5"]]
    # same address as a measurement experience site (2) > same ZIP as one (1) > same ZIP as a condition research site
    docs["exp_score"] = [2.0 if a else (1.0 if "measurement:" in z else (0.5 if z else 0.0))
                         for a, z in zip(exp_addr, exp_zip)]
    docs["experience_match"] = [f"same address: {a}" if a else (f"same ZIP (co-location, not affiliation): {z}"
                                                                 if z else "") for a, z in zip(exp_addr, exp_zip)]
    tiers = specialty_tiers(cond, meas)
    listed = rank_clinicians(docs, tiers, per_county, max_per_address, specialty_before_rx)

    # activity clinicians who are in these counties but outside the listing (NPIs not in NPPES providers, or NPs/PAs)
    billing_in_county = act[act["act_county"].isin(fips) & act["best_code_role"].ne("none")]
    cinfo = counties.set_index("county_fips")
    out_rows = []
    for r in listed.itertuples():
        nm = (f"{str(r.provider_first_name).title()} {str(r.provider_last_name).title()}"
              if isinstance(r.provider_last_name, str) and r.provider_last_name else r.name)
        out_rows.append({
            "County rank": int(cinfo.loc[r.county_fips, "rank"]), "County": cinfo.loc[r.county_fips, "geo_name"],
            "County FIPS": r.county_fips, "Priority in county": int(r.priority), "Name": nm,
            "Credential": r.credential or "", "NPPES specialty": r.taxonomy,
            "Specialty tier": {3: "primary-measurement implementer", 2: "measurement implementer",
                               1: "primary care / condition specialty"}.get(int(r.spec_score), "other"),
            "Measurement activity tier": r.best_code_role if isinstance(r.best_code_role, str) else "none",
            "Measurement activity (codes:services)": r.codes_billed if isinstance(r.codes_billed, str) else "",
            "Medicare services (measurement codes)": int(r.act_services_all) if pd.notna(r.act_services_all) else 0,
            "Part D signals": r.rx_signals if isinstance(r.rx_signals, str) else "",
            "Experience-site match": r.experience_match,
            "Practice address": str(r.address_line1 or "").title(), "City": str(r.city or "").title(),
            "State": r.state, "ZIP": r.zip5, "Phone": _fmt_phone(r.phone), "NPI": str(r.npi),
            "NPI registry": "lookup",
            "NPPES last updated": r.last_update_date.date() if pd.notna(r.last_update_date) else None,
            "Outreach status": "", "Date contacted": None, "Contact person / email": "", "Notes": ""})
    clin = pd.DataFrame(out_rows)
    if len(clin):
        clin = clin.sort_values(["County rank", "Priority in county"], kind="stable").reset_index(drop=True)

    # counties tab
    ctab = counties.copy()
    cap = None
    if table_exists("geo_measurement_capacity"):
        cap = read_table("geo_measurement_capacity")
        cap = cap[(cap["measurement_id"] == mid) & (cap["geo_level"] == "county")].set_index("geo_id")
    lc = clin.groupby("County FIPS") if len(clin) else None
    rows = []
    for r in ctab.itertuples():
        row = {"Rank": int(r.rank), "County": r.geo_name, "State": r.state_abbr, "County FIPS": r.county_fips,
               "Population": int(r.population_total) if pd.notna(r.population_total) else None}
        for c in ctab.columns:
            if c not in ("rank", "county_fips", "geo_name", "state_abbr", "population_total"):
                v = getattr(r, c)
                row[c] = round(float(v), 4) if isinstance(v, (float, np.floating)) and pd.notna(v) else v
        if cap is not None and r.county_fips in cap.index:
            k = cap.loc[r.county_fips]
            row.update({"Capacity basis": k["capacity_basis"],
                        "Billing clinicians in county (Medicare FFS)": k["active_clinicians_n"],
                        "Billing clinicians per 100k adults": round(float(k["active_clinicians_per_100k_adults"]), 2)
                        if pd.notna(k["active_clinicians_per_100k_adults"]) else None,
                        "Billing clinicians within 50 km": k["active_clinicians_within_50km_n"],
                        "Measurement capacity pct": round(float(k["measurement_capacity_pct"]), 3)
                        if pd.notna(k["measurement_capacity_pct"]) else None})
        sub = docs[docs["county_fips"] == r.county_fips]
        row["Eligible clinicians (NPPES, filtered)"] = int(len(sub))
        row["Billing clinicians not listable (not in NPPES providers or not MD/DO)"] = int(
            len(set(billing_in_county.loc[billing_in_county["act_county"] == r.county_fips, "npi"])
                - set(sub["npi"])))
        g = lc.get_group(r.county_fips) if lc is not None and r.county_fips in lc.groups else pd.DataFrame()
        row["Listed"] = int(len(g))
        row["Listed with measurement activity"] = int((g.get("Measurement activity tier", pd.Series(dtype=str))
                                                       != "none").sum()) if len(g) else 0
        row["Listed with dedicated-code activity"] = int((g["Measurement activity tier"] == "dedicated").sum()) \
            if len(g) else 0
        row["Listed with Part D signal"] = int((g["Part D signals"] != "").sum()) if len(g) else 0
        row["Listed with experience-site match"] = int((g["Experience-site match"] != "").sum()) if len(g) else 0
        rows.append(row)
    ctab = pd.DataFrame(rows)

    # sites tab: measurement experience sites + condition research sites in the selected counties
    st = []
    if table_exists("measurement_experience_sites"):
        e = read_table("measurement_experience_sites")
        e = e[(e["measurement_id"] == mid) & e["county_fips"].isin(fips)]
        for r in e.itertuples():
            st.append({"County FIPS": r.county_fips, "Kind": "measurement experience", "Facility": r.facility_name,
                       "Facility type": r.primary_kind, "Address": str(r.address).title() if pd.notna(r.address) else "",
                       "City": str(r.city).title() if pd.notna(r.city) else "", "State": r.state, "ZIP": r.zip5,
                       "Classes matched": r.member_classes_matched, "Conditions": r.condition_ids,
                       "Trials (NCT)": r.nct_ids if pd.notna(r.nct_ids) else "",
                       "NIH projects (appl_id)": r.appl_ids if pd.notna(r.appl_ids) else "",
                       "NIH contact PIs": r.nih_contact_pis if pd.notna(r.nih_contact_pis) else "",
                       "Object id": r.facility_object_id})
    rs = read_table("research_site_registry", columns=["object_id", "facility_name", "primary_kind", "address",
                                                       "city", "state", "zip5", "county_fips",
                                                       "condition_ids_literal", "nih_condition_ids_precision",
                                                       "trial_object_ids", "nih_object_ids"])
    rs = rs[rs["county_fips"].isin(fips)]
    hit = rs["condition_ids_literal"].fillna("").str.split(SEP).map(lambda x: bool(set(x) & set(cond["members"]))) | \
        rs["nih_condition_ids_precision"].fillna("").str.split(SEP).map(lambda x: bool(set(x) & set(cond["members"])))
    for r in rs[hit].itertuples():
        st.append({"County FIPS": r.county_fips, "Kind": "condition research site", "Facility": r.facility_name,
                   "Facility type": r.primary_kind, "Address": str(r.address).title() if pd.notna(r.address) else "",
                   "City": str(r.city).title() if pd.notna(r.city) else "", "State": r.state, "ZIP": r.zip5,
                   "Classes matched": "",
                   "Conditions": SEP.join(sorted({x for x in f"{r.condition_ids_literal}|{r.nih_condition_ids_precision}"
                                                  .split(SEP) if x and x not in ("None", "nan")})),
                   "Trials (NCT)": str(r.trial_object_ids or "").replace("trial:", ""),
                   "NIH projects (appl_id)": str(r.nih_object_ids or "").replace("nih:", ""), "NIH contact PIs": "",
                   "Object id": r.object_id})
    sites = pd.DataFrame(st, columns=["County FIPS", "Kind", "Facility", "Facility type", "Address", "City", "State",
                                      "ZIP", "Classes matched", "Conditions", "Trials (NCT)",
                                      "NIH projects (appl_id)", "NIH contact PIs", "Object id"])
    if len(sites):
        rk = counties.set_index("county_fips")["rank"]
        sites.insert(0, "County rank", sites["County FIPS"].map(rk).astype(int))
        sites = sites.sort_values(["County rank", "Kind", "Facility"])

    n = len(clin)
    n_act = int((clin["Measurement activity tier"] != "none").sum()) if n else 0
    n_ded = int((clin["Measurement activity tier"] == "dedicated").sum()) if n else 0
    n_rx = int((clin["Part D signals"] != "").sum()) if n else 0
    n_exp = int((clin["Experience-site match"] != "").sum()) if n else 0
    n_exp_addr = int(clin["Experience-site match"].str.startswith("same address").sum()) if n else 0
    year = int(act["data_year"].iloc[0]) if "data_year" in act and len(act) else UNKNOWN
    from .facilities.activity import capacity_basis
    cb = capacity_basis(mid)
    member_basis = {m: capacity_basis(m) for m in meas["members"]}
    readme = _readme(cond, meas, basis, weight_set, top_n_counties, per_county, max_per_address,
                     include_nonphysicians, cb, member_basis, year, n, n_act, n_ded, n_rx, n_exp)
    if specialty_before_rx:
        i = next(k for k, t in enumerate(readme) if t.startswith("Ranking within a county"))
        readme.insert(i + 1, "ORDER OVERRIDE (--specialty-before-rx): the specialty tier (4) is applied before Part D "
                             "signals (2) and experience sites (3).")
    out_dir = Path(out_dir)
    if not out_dir.is_absolute():
        out_dir = PROJECT_ROOT / out_dir
    tag = ("_specialty_first" if specialty_before_rx else "") + ("_subgroup" if basis["basis"] == "geo_subgroup_burden"
                                                                  else "")
    path = out_dir / f"outreach_{cond['condition_id'].replace('+', '_or_')}_{mid}_top{top_n_counties}{tag}_" \
                     f"{dt.date.today():%Y-%m-%d}.xlsx"
    _write_workbook(path, readme, ctab, clin, sites)
    return {"path": str(path), "condition_id": cond["condition_id"], "measurement_id": mid,
            "county_basis": basis["basis"], "counties": counties["county_fips"].tolist(),
            "clinicians_listed": n, "with_measurement_activity": n_act, "with_dedicated_code_activity": n_ded,
            "with_partd_signal": n_rx, "with_experience_site_match": n_exp,
            "with_experience_site_same_address": n_exp_addr, "capacity_basis": cb,
            "listed_by_activity_tier": clin["Measurement activity tier"].value_counts().to_dict() if n else {}}


def _readme(cond, meas, basis, ws, top_n, per_county, max_addr, nonphys, cb, member_basis, year, n, n_act, n_ded,
            n_rx, n_exp):
    hc = load_config("measurement_hcpcs")
    notes = []
    for m in meas["members"]:
        spec = (hc.get("measurements") or {}).get(m) or {}
        if spec.get("summary"):
            notes.append(f"  {m} (codes: {member_basis.get(m, 'none')}): {spec['summary']}")
    bnote = ((hc.get("bundles") or {}).get(meas["measurement_id"]) or {}).get("note")
    if bnote:
        notes.append(f"  Bundle: {bnote}")
    no_code = [m for m, b in member_basis.items() if b in ("proxy", "none")]
    return [
        f"Measure It to Cure It: candidate clinician outreach list — {cond['label']} x {meas['label']}",
        f"Built {dt.date.today():%Y-%m-%d} by measure_it.outreach (python -m measure_it.outreach --condition "
        f"{cond['condition_id']} --measurement {meas['measurement_id']} --top {top_n} --per-county {per_county}). "
        "Contains named clinicians from public registries: keep this file out of version control.",
        "",
        "WHAT THIS IS",
        "Candidate clinicians to contact about a pilot of the measurement in candidate counties. It is a starting list "
        "for outreach, not an assessment of any clinician and not a list of who treats the condition.",
        "",
        "COUNTIES",
        f"Selection basis: {basis['basis']}. {basis['description']}. {basis.get('universe', '')}",
        f"Weight set: {ws}." if basis["basis"] == "deployment_opportunities" else
        "Weight set not applicable (burden-only fallback).",
        "",
        "CLINICIANS",
        "Source: CMS NPPES (individual NPIs, active, practice ZIP in the county via ZIP -> ZCTA -> 2024 county). "
        + ("MD/DO plus NPs/PAs." if nonphys else "MD/DO physicians only (rerun with --include-nonphysicians for "
                                                 "NPs/PAs).") + " Pediatric/adolescent taxonomies excluded by default.",
        f"Ranking within a county (up to {per_county}; at most {max_addr} per practice address): (1) Medicare FFS "
        f"billing of the measurement's HCPCS/CPT codes in {year} (dedicated code > related code of medium/high "
        "confidence, then service count; low-confidence related or proxy codes are shown but rank only after "
        "specialty, as they are not measurement-specific); (2) Medicare Part D prescribing of condition-signal drugs "
        "(configs/condition_rx_signals.yaml; direct > related phenotype, then claims; needs >= 2 signal drugs or one "
        "of medium specificity, a single low-specificity drug such as midodrine alone is shown but ranks after "
        "specialty); (3) experience-site match (same practice address as a facility "
        "with a registered trial / NIH project mentioning the measurement > same ZIP as such a facility > same ZIP as a "
        "condition research site); (4) specialty tier (specialist implementer groups of the primary member class > "
        "of other member classes > primary care / the "
        "condition's core specialties, configs/relevance.yaml; curated assumptions); (5) most recently updated NPPES "
        "record. Clinicians with no evidence and no relevant specialty are not listed.",
        f"Claims basis of this measurement: {cb} (dedicated code > related code > labelled proxy > none). "
        + (f"No CPT/HCPCS code exists for {', '.join(no_code)}: billing evidence for it is a PROXY or absent, never the "
           "measurement itself. " if no_code else "")
        + ("No claims code signals this measurement; activity columns are empty by design." if cb == "none" else ""),
        *notes,
        f"In this workbook: {n} clinicians listed; {n_act} with measurement-specific Medicare billing evidence "
        f"({n_ded} on a dedicated code); {n_rx} with a Part D signal; {n_exp} with an experience-site match.",
        "",
        "CAVEATS (read before contacting anyone)",
        "Billing a code does NOT mean the clinician sees Long COVID, ME/CFS, Lyme disease or any target condition; it "
        "means they performed or interpreted that procedure for Medicare fee-for-service patients. A Part D "
        "prescription is not a diagnosis: the configured drugs are not specific to these conditions.",
        "Medicare data cover fee-for-service Part B / Part D only (no Medicare Advantage, Medicaid, commercial or VA "
        "care) and CMS suppresses rows with fewer than 11 beneficiaries (Part D: fewer than 11 claims): absence of "
        "evidence is not evidence of absence, especially for clinicians who mostly see younger patients.",
        "An NPPES taxonomy is self-reported. Addresses and phones are self-reported and may be billing or "
        "administrative; an old 'NPPES last updated' date is more likely stale. Verify via the NPI registry link.",
        "Experience-site 'same ZIP' matches are co-location, not affiliation. A registered study mentioning the "
        "measurement is research experience, not evidence the site offers the test or that the test works.",
        "Counties are candidate deployment opportunities, not a validated ranking (see results/SCORING_RESULTS.md for "
        "rank intervals); no county prevalence is inferred from a state value.",
        "NPPES has no e-mail addresses. Phone numbers are practice numbers.",
        "",
        "TRACKING",
        "Yellow columns on the Clinicians tab (Outreach status, Date contacted, Contact person / email, Notes) are for "
        "you to fill in. Outreach status has a drop-down.",
    ]


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description="Candidate clinician outreach workbook (writes under outreach/)")
    ap.add_argument("--condition", required=True, help="condition id, scored set id, 'a+b' or 'a_or_b'")
    ap.add_argument("--measurement", required=True, help="measurement class or bundle id")
    ap.add_argument("--top", type=int, default=10, help="top-N counties")
    ap.add_argument("--per-county", type=int, default=20)
    ap.add_argument("--weight-set", default="equal")
    ap.add_argument("--out-dir", default="outreach")
    ap.add_argument("--include-nonphysicians", action="store_true", help="also list NPs and PAs")
    ap.add_argument("--max-per-address", type=int, default=5)
    ap.add_argument("--include-pediatric", action="store_true", help="keep pediatric/adolescent taxonomies")
    ap.add_argument("--specialty-before-rx", action="store_true",
                    help="rank the measurement's specialist implementer groups before Part D signals and sites")
    ap.add_argument("--prefer-subgroup", action="store_true",
                    help="rank counties by geo_subgroup_burden (stage E) even when a scored row exists")
    ap.add_argument("--phenotype-id", default=None, help="geo_subgroup_burden phenotype to use")
    a = ap.parse_args(argv)
    res = build_outreach(a.condition, a.measurement, a.top, a.per_county, a.weight_set, a.out_dir,
                         a.include_nonphysicians, a.max_per_address, a.include_pediatric, a.specialty_before_rx,
                         a.phenotype_id, a.prefer_subgroup)
    print(json.dumps(res, indent=2, default=str))


if __name__ == "__main__":
    main()
